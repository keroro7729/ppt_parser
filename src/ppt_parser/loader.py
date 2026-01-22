from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_INCH = 914400
PX_PER_INCH = 96


def emu_to_px(emu):
    if emu is None:
        return None
    return round(emu / EMU_PER_INCH * PX_PER_INCH, 2)


SHAPE_TYPE_MAP = {
    MSO_SHAPE_TYPE.AUTO_SHAPE: "auto_shape",
    MSO_SHAPE_TYPE.TEXT_BOX: "text_box",
    MSO_SHAPE_TYPE.PICTURE: "picture",
    MSO_SHAPE_TYPE.TABLE: "table",
    MSO_SHAPE_TYPE.GROUP: "group",
    MSO_SHAPE_TYPE.PLACEHOLDER: "placeholder",
    MSO_SHAPE_TYPE.FREEFORM: "freeform",
    MSO_SHAPE_TYPE.LINE: "line",
}


# ---------------------------
# 텍스트 (순수 텍스트만)
# ---------------------------
def extract_text(shape):
    if not shape.has_text_frame:
        return None

    texts = []
    for p in shape.text_frame.paragraphs:
        if p.text.strip():
            texts.append(p.text.strip())

    if not texts:
        return None

    return "\n".join(texts)


# ---------------------------
# 테이블 (cell 텍스트만)
# ---------------------------
def extract_table(shape):
    table = shape.table
    rows = []

    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])

    return rows


# ---------------------------
# 단일 Shape → row
# ---------------------------
def extract_shape_row(slide_index, shape_index, shape):
    return {
        "slide_index": slide_index,
        "shape_index": shape_index,
        "shape_type": SHAPE_TYPE_MAP.get(shape.shape_type, "unknown"),

        "left": emu_to_px(shape.left),
        "top": emu_to_px(shape.top),
        "width": emu_to_px(shape.width),
        "height": emu_to_px(shape.height),
        "rotation": shape.rotation,

        "text": extract_text(shape),
        "table": extract_table(shape) if shape.has_table else None,
        "is_line": shape.shape_type == MSO_SHAPE_TYPE.LINE,
    }


# ---------------------------
# PPT 전체 로드
# ---------------------------
def load_pptx(path: str):
    prs = Presentation(path)
    rows = []

    for slide_index, slide in enumerate(prs.slides):
        for shape_index, shape in enumerate(slide.shapes):
            rows.append(
                extract_shape_row(slide_index, shape_index, shape)
            )

    return rows
